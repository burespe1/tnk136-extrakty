#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Převod Wordu do Markdownu (pandoc + úpravy pro extrakty TNK 136).

Zdrojem pravdy je od listopadu 2026 Markdown v extrakty/<WG>/. Tento modul
používá import připomínek (skripty/pripominky.py); samostatně poslouží pro
převod nového extraktu napsaného ve Wordu:

    python skripty/docx2md.py NOVY.docx extrakty/TC278WG4/12345_2025.md

Výsledek je surový převod – hlavičku s metadaty (--- … ---) je třeba doplnit
podle ostatních extraktů. Vyžaduje pandoc; staré .doc převede LibreOffice.
"""

from __future__ import annotations

import argparse
import html
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

KOREN = Path(__file__).resolve().parent.parent

TITULEK = "EXTRAKT z české technické normy"
OBRAZKOVE_PRIPONY = {".png", ".jpg", ".jpeg", ".gif", ".emf", ".wmf", ".bmp", ".tif", ".tiff"}


# --------------------------------------------------------------------------
# konverze
# --------------------------------------------------------------------------

def doc_na_docx(zdroj: Path, kam: Path) -> Path | None:
    """Starý binární .doc převede přes LibreOffice na .docx."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    subprocess.run(
        [soffice, "--headless", "--convert-to", "docx", "--outdir", str(kam), str(zdroj)],
        check=True, capture_output=True, timeout=300,
    )
    vysledek = kam / (zdroj.stem + ".docx")
    return vysledek if vysledek.exists() else None


def titulek_z_docx(docx: Path) -> str:
    """Titulek extraktu je ve Wordu v textovém poli, pandoc ho zahazuje.

    Vytáhneme ho proto přímo z dokumentu. Liší se podle druhu dokumentu
    ("z české technické normy" vs. "z technické specifikace ISO").
    """
    try:
        with zipfile.ZipFile(docx) as z:
            xml = z.read("word/document.xml").decode("utf-8", "replace")
    except Exception:
        return TITULEK
    for odstavec in re.findall(r"<w:p[ >].*?</w:p>", xml, re.S):
        text = "".join(re.findall(r"<w:t[^>]*>(.*?)</w:t>", odstavec, re.S))
        text = " ".join(html.unescape(re.sub(r"<[^>]+>", "", text)).split())
        # text ze tvarů Wordu občas nese na začátku souřadnice obrazce
        m = re.search(r"EXTRAKT\s+z\s+.{3,80}", text)
        if m:
            return m.group(0).strip()
    return TITULEK


def pandoc(docx: Path, media_dir: Path) -> str:
    hotovo = subprocess.run(
        [
            "pandoc", str(docx),
            "--from", "docx",
            "--to", "gfm",
            "--wrap", "none",
            "--extract-media", str(media_dir),
        ],
        check=True, capture_output=True, timeout=300,
    )
    return hotovo.stdout.decode("utf-8")


# --------------------------------------------------------------------------
# úpravy vygenerovaného Markdownu
# --------------------------------------------------------------------------

def _bez_znacek(s: str) -> str:
    s = re.sub(r"<[^>]+>", "\n", s)
    return html.unescape(s)


def hlavicka(text: str) -> str:
    """Nahradí hlavičkovou tabulku extraktu čitelnými řádky.

    Word má v hlavičce jednořádkovou tabulku: název normy | označení + třídicí
    znak. Pandoc z ní udělá buď syrové HTML, nebo rozbitou pipe tabulku, kde
    označení a třídicí znak splynou ("ČSN EN ISO 14819-101 8253"). Obojí
    převedeme na název tučně a pod ním označení s třídicím znakem.
    """

    def slozeni(nazev: str, zbytek: list[str]) -> str:
        radky = []
        if nazev:
            radky += [f"**{nazev}**", ""]
        oznaceni = tridici = ""
        slepene = " ".join(z for z in zbytek if z)
        m = re.match(r"^(?P<ozn>.*?)\s*(?P<tz>\d{2}\s?\d{4})\s*$", slepene)
        if m:
            oznaceni, tridici = m.group("ozn").strip(), m.group("tz").strip()
        else:
            oznaceni = slepene.strip()
        if oznaceni and tridici:
            radky += [f"{oznaceni} — třídicí znak {tridici}", ""]
        elif oznaceni:
            radky += [oznaceni, ""]
        return "\n".join(radky)

    # a) syrové HTML z pandoku
    def z_html(m: re.Match) -> str:
        bunky = [
            " ".join(_bez_znacek(b).split())
            for b in re.findall(r"<td[^>]*>(.*?)</td>", m.group(0), re.S)
        ]
        if not bunky:
            return m.group(0)
        # druhá buňka má označení a třídicí znak ve dvou odstavcích
        syrove = re.findall(r"<td[^>]*>(.*?)</td>", m.group(0), re.S)
        zbytek: list[str] = []
        for b in syrove[1:]:
            zbytek += [" ".join(_bez_znacek(x).split()) for x in re.split(r"</p>", b)]
        return slozeni(bunky[0], zbytek)

    # hlavička končí řádkem „Vydána…“ nebo prvním nadpisem – tabulky za ní nejsou hlavičkové
    m_konec = re.search(r"^(?:Vydána|Platí od|#+ (?!.*(?:EXTRAKT|Extrakt nenahrazuje))|\*\*Úvod)", text, re.M)
    konec = m_konec.start() if m_konec else len(text)
    zahlavi, telo = text[:konec], text[konec:]

    nove = re.sub(r"<table>.*?</table>", z_html, zahlavi, count=1, flags=re.S)
    if nove != zahlavi:
        return nove + telo

    # b) pipe tabulka; řádek s tabulkou může být prázdný (jen „|  |  |“), obsah je pak až pod oddělovačem
    radky = zahlavi.split("\n")
    for i, radek in enumerate(radky):
        if not radek.startswith("|") or not re.search(r"ČSN|ISO|EN |CEN|TNI", radek):
            continue
        j = i + 1
        if i >= 2 and set(radky[i - 1].replace("|", "").replace(" ", "")) <= {"-", ":"} and radky[i - 1].strip():
            od = i - 2  # prázdná hlavička tabulky + oddělovač nad řádkem s obsahem
        elif j < len(radky) and set(radky[j].replace("|", "").replace(" ", "")) <= {"-", ":"} and radky[j].strip():
            od, j = i, i + 1
        else:
            continue
        bunky = [b.strip() for b in radek.strip().strip("|").split("|")]
        bunky = [re.sub(r"^\*\*(.*)\*\*$", r"\1", b) for b in bunky]
        nahrada = slozeni(bunky[0] if bunky else "", bunky[1:])
        return "\n".join(radky[:od] + nahrada.split("\n") + radky[max(i, j) + 1:]) + telo
    return text


def bez_zdvojeneho_titulku(text: str, titulek: str) -> str:
    """Titulek bývá v dokumentu i jako tučný odstavec; necháme jen nadpis."""
    obr = r"(?:!\[[^\]]*\]\([^)]*\))?"
    vzor = re.compile(
        r"^" + obr + r"\s*\*\*" + re.escape(titulek) + r"\*\*\s*" + obr + r"\s*$",
        re.M,
    )
    # logo z textového pole s titulkem se tím zahodí taky – nic na něj neodkazuje
    return vzor.sub("", text, count=1)


def zanor_nadpisy(text: str) -> str:
    """Posune nadpisy o úroveň níž, ať je titulek extraktu jediná jednička."""
    return re.sub(r"^(#{1,5})(?= )", r"#\1", text, flags=re.M)


def popisky_obrazku(text: str) -> str:
    """Popisek pod obrázkem přenese do alternativního textu a sjednotí jeho zápis.

    Ve Wordu je popisek samostatný odstavec ("Obrázek 3 - Ukázka metody ...").
    V Markdownu z něj uděláme alt text obrázku a řádek "**Obr. 3** Ukázka ...",
    aby byl obrázek popsaný i tam, kde se nenačte.
    """
    vzor = re.compile(
        r"^(?P<pred>>[ \t]*)?(?P<img>!\[\]\((?P<src>[^)]+)\))[ \t]*\n"
        r"\n"
        r"(?P<druh>Obrázek|Obr\.|Tabulka|Tab\.)[ \t]*(?P<cislo>\d+)[ \t]*[-–—:.]?[ \t]*(?P<popis>[^\n]+)$",
        re.M,
    )

    def nahrad(m: "re.Match") -> str:
        popis = m.group("popis").strip()
        zkratka = "Tab." if m.group("druh").startswith(("Tabulka", "Tab")) else "Obr."
        alt = popis.replace("[", "(").replace("]", ")")
        return (
            f"{m.group('pred') or ''}![{alt}]({m.group('src')})\n"
            f"**{zkratka} {m.group('cislo')}** {popis}"
        )

    return vzor.sub(nahrad, text)


POPISEK = re.compile(
    r"^(?P<druh>Obrázek|Obr\.|Tabulka|Tabulky|Tab\.|Výpis)[ \t]*(?P<cislo>\d+)\.?"
    r"(?P<oddel>[ \t]*[-–—:][ \t]*|[ \t]+)(?P<popis>\S.*)$"
)


def popisky_samostatne(text: str) -> str:
    """Popisky, které nestojí pod obrázkem (pod tabulkou, ukázkou kódu), sjednotí na **Obr. N** / **Tab. N**.

    Bere jen celý řádek: buď s pomlčkou za číslem („Tabulka 4 - Příklad …“),
    nebo celý tučný („**Tabulka 1 (ukázka z tab. 2) …**“) – věta „Tabulka 3
    obsahuje …“ v běžném textu tak zůstane beze změny.
    """
    def radek(r: str) -> str:
        if r.startswith(("**Obr.", "**Tab.", "**Výpis")) and re.match(r"^\*\*(Obr\.|Tab\.|Výpis) \d+\*\* ", r):
            return r                                   # už sjednocený
        cely_tucny = r.startswith("**") and r.rstrip().endswith("**")
        holy = r.replace("**", "").strip()
        m = POPISEK.match(holy)
        if not m:
            return r
        s_pomlckou = bool(re.search(r"[-–—:]", m.group("oddel")))
        velke = m.group("popis")[:1].isupper()        # „Obrázek 2 Příklad …“ ano, „Obrázek 1 ilustruje …“ ne
        if not (cely_tucny or s_pomlckou or velke):
            return r
        if len(holy) > 300:
            return r
        zkratka = "Tab." if m.group("druh").startswith("Tab") else ("Výpis" if m.group("druh") == "Výpis" else "Obr.")
        return f"**{zkratka} {m.group('cislo')}** {m.group('popis').strip()}"
    text = "\n".join(radek(r) for r in text.split("\n"))
    # alternativní text z Wordu typu „Popis: Clipboard02.jpg“ nebo „FPI Message.emf“ nic neříká
    text = re.sub(r"!\[(?:Popis:[^\]]*|[^\]]*\.(?:png|jpe?g|gif|emf|wmf|bmp|tiff?))\]\(", "![](", text, flags=re.I)
    alt = lambda s: s.strip().replace("[", "(").replace("]", ")")
    # popisek obrázku NAD obrázkem (ve Wordu běžné u výřezů tabulek) → pod obrázek
    text = re.sub(
        r"^(?P<pop>\*\*Obr\. \d+\*\* (?P<popis>[^\n]+))\n(?:[ \t]*\n){0,2}!\[\]\((?P<src>[^)]+)\)[ \t]*$",
        lambda m: f"![{alt(m.group('popis'))}]({m.group('src')})\n{m.group('pop')}",
        text, flags=re.M,
    )
    # popisek tabulky nad výřezem tabulky: zůstane nad, obrázek dostane alternativní text
    text = re.sub(
        r"^(?P<pop>\*\*Tab\. \d+\*\* (?P<popis>[^\n]+))(?P<mezi>\n(?:[ \t]*\n){0,2})!\[\]\((?P<src>[^)]+)\)",
        lambda m: f"{m.group('pop')}{m.group('mezi')}![{alt(m.group('popis'))}]({m.group('src')})",
        text, flags=re.M,
    )
    # obrázek bez alternativního textu těsně nad popiskem dostane text z popisku
    return re.sub(
        r"!\[\]\((?P<src>[^)]+)\)(?P<mezi>[ \t]*\n(?:[ \t]*\n){0,2})(?P<pop>\*\*Obr\. \d+\*\* (?P<popis>[^\n]+))",
        lambda m: f"![{alt(m.group('popis'))}]({m.group('src')}){m.group('mezi')}{m.group('pop')}",
        text,
    )


PRILOHA = re.compile(r"^(?:\*\*)?(Příloha\s+[A-Z]\b[^*\n]{0,200}?)(?:\*\*)?\s*$")


def nadpisy_priloh(text: str) -> str:
    """„Příloha A (normativní) …“ jako samostatný (tučný) odstavec převede na nadpis."""
    def radek(r: str) -> str:
        if r.startswith("#"):
            return r
        cisty = r.replace("**", "")
        cisty = re.sub(r"(?<![\w\\])_(.+?)_(?!\w)", r"\1", cisty)        # kurzíva _x_
        cisty = re.sub(r"(?<![\w*\\])\*(.+?)\*(?!\w)", r"\1", cisty).strip()  # kurzíva *x* (pandoc)
        cely_tucny = r.startswith("**") and r.rstrip().endswith("**")
        m = PRILOHA.match(cisty)
        if not m:
            return r
        obsah = m.group(1).strip()
        # musí to být nadpis, ne věta: buď (normativní)/(informativní), nebo celý tučný řádek
        if not (re.search(r"\((normativní|informativní)\)", obsah) or cely_tucny):
            return r
        return f"## {obsah}"
    return "\n".join(radek(r) for r in text.split("\n"))


JEDNA_BUNKA = re.compile(
    r"<table>\s*(?:<colgroup>.*?</colgroup>\s*)?<tbody>\s*<tr[^>]*>\s*<td[^>]*>(?P<obsah>.*?)</td>\s*</tr>\s*</tbody>\s*</table>",
    re.S,
)


def kod_z_tabulek(text: str) -> str:
    """Tabulka s jedinou buňkou, ve které je XML nebo pseudokód, se převede na blok kódu."""
    def nahrad(m: "re.Match") -> str:
        obsah = m.group("obsah")
        if "<img" in obsah or "<table" in obsah:
            return m.group(0)
        odstavce = re.findall(r"<p>(.*?)</p>", obsah, re.S) or [obsah]
        radky = []
        for o in odstavce:
            o = re.sub(r"<br\s*/?>", "\n", o)
            o = re.sub(r"</?(?:strong|em|span|a|sup|sub|code|u)[^>]*>", "", o)
            radky.extend(html.unescape(o).split("\n"))
        radky = [r.rstrip() for r in radky]
        while radky and not radky[0].strip():
            radky.pop(0)
        while radky and not radky[-1].strip():
            radky.pop()
        if not radky:
            return m.group(0)
        prvni = radky[0].lstrip()
        if prvni.startswith("<"):
            jazyk = "xml"
        elif any(":=" in r for r in radky):
            jazyk = ""
        else:
            return m.group(0)                          # běžná tabulka s textem
        return "```" + jazyk + "\n" + "\n".join(radky) + "\n```"
    return JEDNA_BUNKA.sub(nahrad, text)


SKUPINY = ("wordprocessingGroup", "wordprocessingCanvas")
EMU_NA_PX = 9525                      # 96 dpi


def vykresli_skupiny(zdroj: Path, tmp: Path) -> Path:
    """Kresby, které pandoc vynechá, nahradí v kopii dokumentu obrázkem.

    - skupiny a plátna z tvarů (s textovými poli): každá se vykreslí zvlášť –
      kopie dokumentu jen s jejím odstavcem → PDF (LibreOffice) → PNG (pdftoppm)
    - vložené objekty OLE (Visio, PowerPoint …): použije se jejich náhled
    Bez LibreOffice, pdftoppm nebo python-docx vrátí zdroj beze změny.
    """
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return zdroj
    try:
        import docx                                   # python-docx
        from docx.oxml.ns import qn
        from docx.shared import Emu
        from PIL import Image
    except ImportError:
        return zdroj

    dok = docx.Document(str(zdroj))
    MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
    WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    cile = []                                         # (běh, cx, cy)
    for r in dok.element.body.iter(qn("w:r")):
        if any(a.tag == qn("w:del") for a in r.iterancestors()):
            continue                                  # smazaná revize – pandoc ji zahodí
        for kresba in r.iter(qn("w:drawing")):
            if any(a.tag == MC + "Fallback" for a in kresba.iterancestors()):
                continue
            g = kresba.find(".//" + A + "graphicData")
            ext = kresba.find(".//" + WP + "extent")
            if g is None or ext is None or not g.get("uri", "").endswith(SKUPINY):
                continue
            cx, cy = int(ext.get("cx")), int(ext.get("cy"))
            if cx < 2 * 360000:                        # drobnosti pod 2 cm vynechat
                continue
            cile.append((r, cx, cy))
    # vložené objekty OLE (Visio, PowerPoint …): pandoc je vynechá, Word ukazuje jejich náhled
    V = "{urn:schemas-microsoft-com:vml}"
    R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    ole = 0
    for r in list(dok.element.body.iter(qn("w:r"))):
        obj = r.find(qn("w:object"))
        if obj is None or any(a.tag == qn("w:del") for a in r.iterancestors()):
            continue
        img = obj.find(".//" + V + "imagedata")
        tvar = obj.find(".//" + V + "shape")
        rid = img.get(R + "id") if img is not None else None
        if not rid or rid not in dok.part.related_parts:
            continue
        cast = dok.part.related_parts[rid]
        nahled = tmp / ("ole_%d%s" % (ole, Path(cast.partname).suffix))
        nahled.write_bytes(cast.blob)
        nahled = na_png(nahled)
        if nahled.suffix.lower() != ".png" and nahled.suffix.lower() not in (".jpg", ".jpeg", ".gif"):
            continue
        styl = tvar.get("style", "") if tvar is not None else ""
        w = re.search(r"width:([\d.]+)pt", styl); h = re.search(r"height:([\d.]+)pt", styl)
        p_el = next((a for a in r.iterancestors() if a.tag == qn("w:p")), None)
        if p_el is None:
            continue
        novy = docx.text.paragraph.Paragraph(p_el, dok).add_run()
        if w and h:
            novy.add_picture(str(nahled), width=Emu(int(float(w.group(1)) * 12700)), height=Emu(int(float(h.group(1)) * 12700)))
        else:
            novy.add_picture(str(nahled))
        r.addprevious(novy._r)
        r.getparent().remove(r)
        ole += 1

    if not cile:
        if ole:
            cil = tmp / (zdroj.stem + "_kresby.docx")
            dok.save(str(cil))
            return cil
        return zdroj

    # každou skupinu vykreslit zvlášť: kopie dokumentu jen s jejím odstavcem → PDF → PNG
    pdftoppm = shutil.which("pdftoppm")
    if not pdftoppm:
        cile = []
    telo = dok.element.body
    deti = list(telo.iterchildren())
    nahrazeno = 0
    for n, (beh, cx, cy) in enumerate(cile):
        horni = next((a for a in [beh, *beh.iterancestors()] if a.getparent() is telo), None)
        p_el = next((a for a in beh.iterancestors() if a.tag == qn("w:p")), None)
        if horni is None or p_el is None:
            continue
        idx = deti.index(horni)
        kopie = docx.Document(str(zdroj))
        k_telo = kopie.element.body
        for k, el in enumerate(list(k_telo.iterchildren())):
            if k != idx and el.tag != qn("w:sectPr"):
                k_telo.remove(el)
        for el in list(k_telo.iter(qn("w:del"), qn("w:moveFrom"))):   # revize přijmout,
            el.getparent().remove(el)                                    # jako to dělá pandoc
        for el in list(k_telo.iter(qn("w:ins"), qn("w:moveTo"))):
            for dite in list(el):
                el.addprevious(dite)
            el.getparent().remove(el)
        for sp in k_telo.iter(qn("w:sectPr")):             # bez záhlaví a zápatí
            for ref in sp.findall(qn("w:headerReference")) + sp.findall(qn("w:footerReference")):
                sp.remove(ref)
        jedna = tmp / f"kresba_{n}.docx"
        kopie.save(str(jedna))
        try:
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(jedna)],
                           capture_output=True, timeout=120, check=True)
            subprocess.run([pdftoppm, "-r", "150", "-f", "1", "-l", "1", "-png", "-singlefile",
                            str(jedna.with_suffix(".pdf")), str(jedna.with_suffix(""))],
                           capture_output=True, timeout=120, check=True)
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired):
            continue
        png = jedna.with_suffix(".png")
        if not png.exists():
            continue
        _orizni(png)
        novy = docx.text.paragraph.Paragraph(p_el, dok).add_run()
        novy.add_picture(str(png), width=Emu(cx), height=Emu(cy))
        beh.addprevious(novy._r)
        beh.getparent().remove(beh)
        nahrazeno += 1
    if not nahrazeno and not ole:
        return zdroj
    cil = tmp / (zdroj.stem + "_kresby.docx")
    dok.save(str(cil))
    return cil


XML_RADEK = re.compile(r"^[ \t]*\\<\\?[/?!a-zA-Z]")


def kod_z_odstavcu(text: str) -> str:
    """Ukázka XML zapsaná ve Wordu jako běžné odstavce (pandoc píše „\\<xs:…\\>“) → blok kódu.

    Bere souvislý úsek aspoň dvou řádků začínajících escapovaným „<“; prázdné
    řádky mezi nimi (odstavce ve Wordu) se v bloku vynechají.
    """
    radky = text.split("\n")
    out, i = [], 0
    while i < len(radky):
        if XML_RADEK.match(radky[i]):
            j, blok = i, []
            while j < len(radky) and (XML_RADEK.match(radky[j]) or not radky[j].strip()
                                      or radky[j].strip() in ("…", "...")
                                      or (blok and not blok[-1].rstrip().endswith(">"))):  # pokračování tagu
                if radky[j].strip():
                    blok.append(radky[j])
                j += 1
            if len(blok) >= 2:
                cisty = [re.sub(r"\\([\\`*_{}\[\]()#+\-.!|<>~])", r"\1", r).rstrip() for r in blok]
                out += ["```xml", *cisty, "```", ""]
                i = j
                continue
        out.append(radky[i])
        i += 1
    return "\n".join(out)


# --------------------------------------------------------------------------
# bloky kódu (styl „Source Code“ z md2docx)
# --------------------------------------------------------------------------

ZNACKA_KODU = "KODBLOK{}KONEC"


def vyjmi_kod(docx: Path, tmp: Path) -> "tuple[Path, list[str]]":
    """Odstavce ve stylu Source Code (bloky kódu z md2docx) nahradí v kopii dokumentu značkou.

    Pandoc je z Wordu čte jako běžný text a ztrácí odsazení i prázdné řádky; text
    se proto vezme přímo z XML (w:t, w:tab, w:br) a po převodu vloží zpět jako blok.
    """
    from lxml import etree
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    q = lambda t: "{%s}%s" % (W, t)
    with zipfile.ZipFile(docx) as z:
        doc = etree.fromstring(z.read("word/document.xml"))
        bloky = []
        for p in list(doc.iter(q("p"))):
            st = p.find(q("pPr") + "/" + q("pStyle"))
            if st is None or st.get(q("val")) not in ("SourceCode", "Source Code"):
                continue
            casti = []
            for el in p.iter(q("t"), q("tab"), q("br")):
                if el.tag == q("t"):
                    casti.append(el.text or "")
                else:
                    casti.append("\t" if el.tag == q("tab") else "\n")
            bloky.append("".join(casti).rstrip("\n"))
            novy = etree.SubElement(p.getparent(), q("p"))
            r = etree.SubElement(novy, q("r")); t = etree.SubElement(r, q("t"))
            t.text = ZNACKA_KODU.format(len(bloky) - 1)
            p.addprevious(novy); p.getparent().remove(p)
        if not bloky:
            return docx, []
        cil = tmp / ("kod_" + docx.name)
        with zipfile.ZipFile(cil, "w", zipfile.ZIP_DEFLATED) as out:
            for i in z.infolist():
                data = z.read(i.filename)
                if i.filename == "word/document.xml":
                    data = etree.tostring(doc, xml_declaration=True, encoding="UTF-8", standalone=True)
                out.writestr(i, data)
    return cil, bloky


def jazyk_kodu(kod: str) -> str:
    """Jazyk bloku podle obsahu (když ho nelze převzít ze zdrojového Markdownu)."""
    if re.search(r"<!(ENTITY|ELEMENT|ATTLIST)\b", kod):
        return "dtd"
    if re.search(r"^\s*<(\?xml|xs:|[A-Za-z_][\w.-]*[ >/])", kod, re.M) and "::=" not in kod and ">:=" not in kod:
        return "xml"
    if "::=" in kod and re.search(r"\b(SEQUENCE|CHOICE|INTEGER|BEGIN|DEFINITIONS|BIT STRING|OCTET STRING|ENUMERATED)\b|\{\s*iso\s*\(", kod):
        return "asn1"
    return ""


def vloz_kod(text: str, bloky: "list[str]", jazyky: "list[str] | None" = None) -> str:
    for i, kod in enumerate(bloky):
        jazyk = jazyky[i] if jazyky and i < len(jazyky) else jazyk_kodu(kod)
        text = re.sub(r"^[ \t]*" + ZNACKA_KODU.format(i) + r"[ \t]*$",
                      lambda m: "```" + jazyk + "\n" + kod + "\n```", text, count=1, flags=re.M)
    return text


PREVEST_NA_PNG = {".emf", ".wmf", ".bmp", ".tif", ".tiff"}


def _orizni(png: Path, okraj: int = 8) -> None:
    """Ořízne bílé okraje (LibreOffice kreslí EMF na celou stránku A4)."""
    try:
        from PIL import Image, ImageChops
    except ImportError:
        return
    im = Image.open(png).convert("RGB")
    bila = Image.new("RGB", im.size, (255, 255, 255))
    rozdil = ImageChops.difference(im, bila).convert("L").point(lambda x: 255 if x > 6 else 0)
    box = rozdil.getbbox()
    if not box:
        return
    l, h, p, d = box
    box = (max(0, l - okraj), max(0, h - okraj), min(im.width, p + okraj), min(im.height, d + okraj))
    if box != (0, 0, im.width, im.height):
        im.crop(box).save(png)


def na_png(cesta: Path) -> Path:
    """EMF/WMF a další formáty, které GitHub nezobrazí, převede na PNG.

    Vektorové EMF/WMF jdou přes PDF (LibreOffice) a pdftoppm ve 150 dpi, aby
    diagram zůstal čitelný; bez pdftoppm přímo LibreOffice do PNG. Bílé okraje
    se ořežou. Když převod selže, vrátí původní soubor.
    """
    if cesta.suffix.lower() not in PREVEST_NA_PNG:
        return cesta
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return cesta
    png = cesta.with_suffix(".png")
    pdftoppm = shutil.which("pdftoppm")
    try:
        if pdftoppm and cesta.suffix.lower() in (".emf", ".wmf"):
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(cesta.parent), str(cesta)],
                           capture_output=True, timeout=120, check=True)
            pdf = cesta.with_suffix(".pdf")
            subprocess.run([pdftoppm, "-r", "150", "-png", "-singlefile", str(pdf), str(png.with_suffix(""))],
                           capture_output=True, timeout=120, check=True)
        else:
            subprocess.run([soffice, "--headless", "--convert-to", "png", "--outdir", str(cesta.parent), str(cesta)],
                           capture_output=True, timeout=120, check=True)
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired):
        return cesta
    if not (png.exists() and png.stat().st_size > 0):
        return cesta
    _orizni(png)
    return png


def presun_obrazky(text: str, media_dir: Path, img_dir: Path, zaklad: str) -> str:
    """Média vytažená pandocem přejmenuje na <zaklad>_imageN.<ext> a přesune do img/.

    Číslují se podle pořadí výskytu v textu; obrázky, na které po úklidu nic
    neodkazuje (typicky pozadí textových polí), se nekopírují.
    """
    dostupne = {
        p.name: p
        for p in media_dir.rglob("*")
        if p.is_file() and p.suffix.lower() in OBRAZKOVE_PRIPONY
    }
    if not dostupne:
        return text
    dostupne = {jmeno: na_png(cesta) for jmeno, cesta in dostupne.items()}

    prejmenovani: dict[str, str] = {}

    def oprav(m: "re.Match") -> str:
        jmeno = re.split(r"[\\/]", m.group("cesta"))[-1]
        if jmeno not in dostupne:
            return m.group(0)
        novy = prejmenovani.get(jmeno)
        if novy is None:
            novy = f"{zaklad}_image{len(prejmenovani) + 1}{dostupne[jmeno].suffix.lower()}"
            prejmenovani[jmeno] = novy
        return f"{m.group('pred')}(./img/{novy})"

    text = re.sub(r"(?P<pred>!\[[^\]]*\])\((?P<cesta>[^)]+)\)", oprav, text)

    # obrázky uvnitř tabulek, které pandoc nechal v HTML
    def oprav_html(m: "re.Match") -> str:
        nahrada = oprav(re.match(r"(?P<pred>)\((?P<cesta>[^)]+)\)", f"({m.group('cesta')})"))
        return f'{m.group("pred")}{nahrada[1:-1]}"' if nahrada.startswith("(") else m.group(0)

    text = re.sub(r'(?P<pred><img src=")(?P<cesta>[^"]+)"', oprav_html, text)

    if prejmenovani:
        img_dir.mkdir(parents=True, exist_ok=True)
        for puvodni, novy in prejmenovani.items():
            shutil.copy2(dostupne[puvodni], img_dir / novy)
    return text


def kurziva_podtrzitkem(text: str) -> str:
    """Pandoc píše kurzívu hvězdičkami, v repozitáři se drží podtržítka (mimo bloky kódu)."""
    casti = re.split(r"(^```.*?^```)", text, flags=re.M | re.S)
    return "".join(
        c if c.startswith("```") else re.sub(r"(?<!\*)\*(?!\*)([^*\n]+?)\*(?!\*)", r"_\1_", c)
        for c in casti
    )


def uklid_text(text: str) -> str:
    text = kurziva_podtrzitkem(text)
    text = text.replace(" ", " ")
    text = re.sub(r"^\s*[•●]\s+", "- ", text, flags=re.M)
    # podtržení odkazu ve Wordu převádí pandoc na <span class="underline">, v Markdownu je zbytečné
    text = re.sub(r'<span class="underline">(.*?)</span>', r"\1", text)
    text = re.sub(r"[ \t]+$", "", text, flags=re.M)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip() + "\n"


def prevod(docx: Path, img_dir: Path, zaklad: str, jazyky: "list[str] | None" = None) -> str:
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        zdroj = docx
        if docx.suffix.lower() == ".doc":
            prevedeny = doc_na_docx(docx, tmp_path)
            if prevedeny is None:
                raise RuntimeError("starý formát .doc a není k dispozici LibreOffice")
            zdroj = prevedeny
        zdroj = vykresli_skupiny(zdroj, tmp_path)
        zdroj, bloky_kodu = vyjmi_kod(zdroj, tmp_path)
        media_dir = tmp_path / "media"
        titulek = titulek_z_docx(docx)
        text = pandoc(zdroj, media_dir)
        text = text.replace("\u00a0", " ")  # Word sype nezlomitelné mezery
        text = vloz_kod(text, bloky_kodu, jazyky)
        text = hlavicka(text)
        text = zanor_nadpisy(text)
        text = bez_zdvojeneho_titulku(text, titulek)
        text = popisky_obrazku(text)
        text = kod_z_tabulek(text)
        text = kod_z_odstavcu(text)
        text = popisky_samostatne(text)
        text = nadpisy_priloh(text)
        text = presun_obrazky(text, media_dir, img_dir, zaklad)
        return uklid_text(f"# {titulek}\n\n{text}")


# --------------------------------------------------------------------------

def main() -> int:
    ap = argparse.ArgumentParser(description="Převede Word extraktu do Markdownu (bez metadat).")
    ap.add_argument("docx", type=Path, help="zdrojový .docx (nebo .doc)")
    ap.add_argument("md", type=Path, help="cílový .md; obrázky se uloží do img/ vedle něj")
    args = ap.parse_args()
    if not shutil.which("pandoc"):
        sys.exit("Chybí pandoc. Nainstaluj ho a spusť znovu.")
    if args.md.exists():
        sys.exit(f"{args.md} už existuje – pro úpravy existujícího extraktu použij pripominky.py.")
    text = prevod(args.docx, args.md.parent / "img", args.md.stem)
    args.md.parent.mkdir(parents=True, exist_ok=True)
    args.md.write_text(text, encoding="utf-8")
    print(f"Převedeno {args.docx} → {args.md} (doplň hlavičku s metadaty)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
